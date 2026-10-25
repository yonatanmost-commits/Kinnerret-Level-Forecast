"""The PowerPoint we hand in (originator 2026-10-23): "a clean powerpoint, no narration no timed
animation, by mouse click only"; the background on a master slide with the page colour and the top
line, the bottom blue line gone, the small lake per slide (not on the master); short notes so a
reader understands the slides; the app slide before the centaur.

Built by docs/talk/build_clean.py -> docs/talk/Kinneret_Talk_Submission.pptx.
"""
import re
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parent.parent
DECK = ROOT / "docs" / "talk" / "Kinneret_Talk_Submission.pptx"

pytestmark = pytest.mark.skipif(not DECK.exists(), reason="the clean deck isn't built")


@pytest.fixture(scope="module")
def prs():
    from pptx import Presentation
    return Presentation(DECK)


TALK_TITLE = "ML-based Forecasting of Sea of Galilee Water Levels Across Multiple Prediction Horizons"


def test_a_title_slide_then_twelve_with_the_app_before_the_centaur(prs):
    titles = [s.shapes.title.text for s in prs.slides]
    assert len(titles) == 13
    assert titles[:5] == [TALK_TITLE, "Lake Kinneret", "The gap", "The project", "Centaur"]


def test_the_title_slide_names_both_authors_and_affiliations(prs):
    """Adnan, 25 Oct: "a title slide at the beginning ... our names and affiliation"; his line
    verbatim. The title as the conference programme prints it."""
    cover = prs.slides[0]
    words = "\n".join(sh.text_frame.text for sh in cover.shapes if sh.has_text_frame)
    assert "Yonatan Most" in words
    assert "Adnan Agbaria" in words
    assert "Software Engineering Department, The Kinneret Academic College" in words
    assert "12th Kinneret Conference" in words
    assert "11 November 2026" in words
    assert 'nodeType="clickEffect"' not in cover._element.xml          # nothing to click through


def _entering(slide) -> set[str]:
    """Shape ids that an entrance effect brings in (hidden until then)."""
    ids = set()
    for ctn in slide._element.iter("{http://schemas.openxmlformats.org/presentationml/2006/main}cTn"):
        if ctn.get("presetClass") == "entr":
            tgt = next(ctn.iter("{http://schemas.openxmlformats.org/presentationml/2006/main}spTgt"), None)
            if tgt is not None:
                ids.add(tgt.get("spid"))
    return ids


def test_no_video_poster_shows_before_its_click(prs):
    """Bug, 25 Oct: Morph from the cover into "Lake Kinneret" drew the video's poster (its LAST
    frame) over the first-frame still, then jumped back. A video waits hidden until its click."""
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    for i, s in enumerate(prs.slides, 1):
        entering = _entering(s)
        for sh in s.shapes:
            if sh.shape_type == MSO_SHAPE_TYPE.MEDIA:
                assert str(sh.shape_id) in entering, f"slide {i}: {sh.name} visible before its click"


def test_the_cover_lake_morphs_into_the_small_lake(prs):
    """Originator, 25 Oct: a Kinneret on the title slide, and a nice Morph into the animation slide."""
    cover, first = prs.slides[0], prs.slides[1]
    big = next(sh for sh in cover.shapes if sh.name == "!!lake")
    small = next(sh for sh in first.shapes if sh.name == "!!lake")
    assert big.height > 5 * small.height
    assert "morph" in first._element.xml.lower()


def test_no_sound_and_no_timings(prs):
    for i, s in enumerate(prs.slides, 1):
        x = s._element.xml
        assert "advTm" not in x, f"slide {i} advances on a timer"
        delays = re.findall(r'<p:cTn [^>]*delay="(\d+)"', x)
        assert all(d == "0" for d in delays), f"slide {i}: timed delays {delays}"
        for rel in s.part.rels.values():
            assert "audio" not in rel.reltype, f"slide {i} has sound"
            if "video" in rel.reltype or "media" in rel.reltype:
                assert "mp4" in rel.target_ref


def test_steps_wait_for_a_click(prs):
    for i, s in enumerate(prs.slides, 1):
        x = s._element.xml
        auto = re.findall(r'nodeType="(afterEffect|withEffect)"', x)
        clicks = x.count('nodeType="clickEffect"')
        if auto:
            assert clicks, f"slide {i}: effects run with no click"


def test_the_master_carries_the_page_and_the_top_rule_only(prs):
    m = prs.slide_master
    names = [sh.name for sh in m.shapes]
    assert "title bar" in names
    assert not any("bottom" in n.lower() for n in names)
    assert 'val="FBFAF7"' in m._element.xml                        # the page colour
    for s in list(prs.slides)[1:]:
        assert s.slide_layout.name == "Kinneret title"
        assert "!!lake" in [sh.name for sh in s.shapes]             # the lake stays on the slide


def test_short_reader_notes_on_every_slide(prs):
    for i, s in enumerate(prs.slides, 1):
        notes = s.notes_slide.notes_text_frame.text.strip()
        assert 25 <= len(notes.split()) <= 120, f"slide {i}: {len(notes.split())} words of notes"
        assert "[" not in notes, f"slide {i}: script markers in the notes"
