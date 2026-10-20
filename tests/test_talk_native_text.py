"""Native text in the talk show (originator 2026-10-20): "build the pptx with more native text, so
i can design some of it by hand". Scope agreed: frame text (subtitles, credits, stamps, legends,
footnotes) plus row names, headings and table text; axis ticks, labels pinned to a data point and
numbers inside pills stay in the picture. Hand edits survive a rebuild by name.

Contract fixed here, in docs/talk/:
- theme.Labels: a chart's words that go native. `L.text(ax, name, x, y, words, size, bold, color,
  **kw)` draws like ax.text (px axes, y down) and `L.centered(...)` like theme.text_centered.
  `L.save(fig, path)` writes the full picture at `path` (unchanged for the review page), the
  picture WITHOUT those words at `<dir>/bare/<name>`, and `<dir>/bare/<stem>.labels.json`:
  {"labels": [{"name", "text", "box": [x0, y0, x1, y1] px y-down, "ha", "pt" (slide points),
  "bold", "color", "lines" (line count)}]}.
- build_show.build(out, segs, audio): the show. Each label becomes a text box named after it on
  its talk slide, faded in on the first click whose picture has it and out on the first that
  doesn't. If `out` exists the build updates it: the shapes it made (pictures, sound, frame) are
  rebuilt; a label box keeps the position, size and formatting it has in the file; its words
  follow the chart unless they were changed by hand (then the hand words stay, and a warning is
  returned if the chart's words changed too); shapes the build didn't make stay as they are.
  Returns the list of warnings.
"""
import json
import subprocess
import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parent.parent
TALK = ROOT / "docs" / "talk"
sys.path.insert(0, str(TALK))

SHOW = TALK / "Kinneret_Talk_Show.pptx"


# ---------------------------------------------------------------- theme.Labels

def chart(tmp_path, words="Photo: Someone, CC BY 2.0"):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    import theme as T
    fig = plt.figure(figsize=(T.W / T.DPI, T.H / T.DPI), dpi=T.DPI, facecolor=T.SURFACE)
    ax = fig.add_axes([0, 0, 1, 1])
    ax.set_xlim(0, T.W)
    ax.set_ylim(T.H, 0)
    ax.axis("off")
    ax.plot([200, 1700], [600, 600], color=T.USED, lw=6)          # a mark that stays in the picture
    L = T.Labels()
    L.text(ax, "credit", 90, 1018, words, 18, False, T.INK_2, ha="left", va="baseline")
    L.centered(ax, "head", 960, 300, "Week\nahead", 40, True, T.INK)
    return L.save(fig, tmp_path / "c.png")


def test_labels_writes_full_picture_bare_picture_and_their_words(tmp_path):
    import numpy as np
    from PIL import Image
    full = chart(tmp_path)
    bare = tmp_path / "bare" / "c.png"
    meta = json.loads((tmp_path / "bare" / "c.labels.json").read_text(encoding="utf-8"))
    assert full == tmp_path / "c.png" and full.exists() and bare.exists()
    labels = {d["name"]: d for d in meta["labels"]}
    assert set(labels) == {"credit", "head"}
    c = labels["credit"]
    assert c["text"] == "Photo: Someone, CC BY 2.0" and c["ha"] == "left" and not c["bold"]
    assert c["pt"] == pytest.approx(18 / 1.2, abs=0.01)           # 1920 px = 960 pt
    x0, y0, x1, y1 = c["box"]
    assert 85 <= x0 <= 95 and y1 <= 1030 and y0 >= 990 and x1 > x0 + 300
    assert labels["head"]["lines"] == 2 and labels["head"]["bold"] and labels["head"]["ha"] == "center"
    a, b = (np.asarray(Image.open(p).convert("RGB"), int) for p in (full, bare))
    X0, Y0, X1, Y1 = (round(v) for v in c["box"])
    assert np.abs(a[Y0:Y1, X0:X1] - b[Y0:Y1, X0:X1]).sum() > 0     # the words are gone from the bare one
    assert (np.abs(a - b).sum(axis=2) > 0)[590:610].sum() == 0     # the line is in both


# ---------------------------------------------------------------- build_show.build (PowerPoint)

def silent_mp3(path: Path, secs: float = 1.0) -> None:
    subprocess.run(["ffmpeg", "-y", "-loglevel", "error", "-f", "lavfi", "-i",
                    "anullsrc=r=24000:cl=mono", "-t", str(secs), "-b:a", "32k", str(path)], check=True)


@pytest.fixture
def talk(tmp_path):
    """Two talk slides: 1 has two clicks (the row appears on the second), 2 has one click."""
    import build_show as B
    audio = tmp_path / "audio"
    audio.mkdir()
    pics = {}
    for key, labels in {"a1": {"head": "WHO"}, "a2": {"head": "WHO", "row": "Water Authority"},
                        "b1": {"credit": "Photo: JNF Photo Archive"}}.items():
        pics[key] = picture(tmp_path, key, labels)
    segs = [dict(id="s01_c1", n="1", q="q1", click=1, clicks=2, image=str(pics["a1"]), text="one", pause=0.0,
                 title="The gap"),
            dict(id="s01_c2", n="1", q="q1", click=2, clicks=2, image=str(pics["a2"]), text="two", pause=0.0,
                 title="The gap"),
            dict(id="s02_c1", n="2", q="q2", click=1, clicks=1, image=str(pics["b1"]), text="three", pause=0.0,
                 title="The limiting factor")]
    for g in segs:
        silent_mp3(audio / f"{g['id']}.mp3")
    return B, segs, audio, tmp_path


def picture(tmp_path: Path, key: str, labels: dict) -> Path:
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    import theme as T
    fig = plt.figure(figsize=(T.W / T.DPI, T.H / T.DPI), dpi=T.DPI, facecolor=T.SURFACE)
    ax = fig.add_axes([0, 0, 1, 1])
    ax.set_xlim(0, T.W)
    ax.set_ylim(T.H, 0)
    ax.axis("off")
    L = T.Labels()
    for k, (name, words) in enumerate(labels.items()):
        L.text(ax, name, 120, 300 + 200 * k, words, 32, name == "head", T.INK, ha="left", va="center")
    return L.save(fig, tmp_path / f"{key}.png")


def shapes(out: Path):
    from pptx import Presentation
    return [{sh.name: sh for sh in sl.shapes} for sl in Presentation(out).slides]


def test_labels_become_named_text_boxes(talk):
    B, segs, audio, tmp = talk
    out = tmp / "show.pptx"
    assert B.build(out=out, segs=segs, audio=audio) == []
    s1, s2 = shapes(out)
    assert s1["head"].text_frame.text == "WHO" and s1["row"].text_frame.text == "Water Authority"
    assert s2["credit"].text_frame.text == "Photo: JNF Photo Archive"
    run = s1["head"].text_frame.paragraphs[0].runs[0]
    assert run.font.name == "Segoe UI Semibold" and run.font.size.pt == pytest.approx(32 / 1.2, abs=0.1)
    assert s2["credit"].text_frame.paragraphs[0].runs[0].font.name == "Segoe UI Semilight"
    assert s1["!!title"].text_frame.text == "The gap"
    x0 = json.loads((tmp / "bare" / "a2.labels.json").read_text(encoding="utf-8"))["labels"][1]["box"][0]
    assert s1["row"].left / 12700 == pytest.approx(x0 * 0.5, abs=2)        # EMU -> pt; px * 0.5 = pt
    # the row fades in with the second click; the head is there from the start
    import win32com.client
    app = win32com.client.Dispatch("PowerPoint.Application")
    pres = app.Presentations.Open(str(out), ReadOnly=True, WithWindow=False)
    try:
        seq = pres.Slides(1).TimeLine.MainSequence
        animated = {seq.Item(i).Shape.Name for i in range(1, seq.Count + 1)}
    finally:
        pres.Close()
    assert "row" in animated and "head" not in animated


def test_rebuild_keeps_hand_design_and_hand_words(talk):
    import win32com.client
    B, segs, audio, tmp = talk
    out = tmp / "show.pptx"
    B.build(out=out, segs=segs, audio=audio)
    pics_before = [sum(sh.shape_type == 13 for sh in s.values()) for s in shapes(out)]
    app = win32com.client.Dispatch("PowerPoint.Application")
    pres = app.Presentations.Open(str(out), WithWindow=False)
    head = pres.Slides(1).Shapes("head")
    head.Left, head.Top = 300, 200                                # moved and restyled by hand
    head.TextFrame.TextRange.Font.Size = 44
    head.TextFrame.TextRange.Font.Color.RGB = 0x0C41C2
    pres.Slides(2).Shapes("credit").TextFrame.TextRange.Text = "Photo: JNF archive (hand)"
    note = pres.Slides(2).Shapes.AddTextbox(1, 100, 100, 200, 40)
    note.Name = "my note"
    note.TextFrame.TextRange.Text = "mine"
    pres.Save()
    pres.Close()
    # the charts change: the row and the credit get new words
    segs[1]["image"] = str(picture(tmp, "a2", {"head": "WHO", "row": "Israel Water Authority"}))
    segs[2]["image"] = str(picture(tmp, "b1", {"credit": "Photo: KKL-JNF Photo Archive"}))
    warnings = B.build(out=out, segs=segs, audio=audio)
    s1, s2 = shapes(out)
    h = s1["head"]
    assert (h.left / 12700, h.top / 12700) == pytest.approx((300, 200), abs=0.5)
    run = h.text_frame.paragraphs[0].runs[0]
    assert run.font.size.pt == 44 and str(run.font.color.rgb) == "C2410C"
    assert s1["row"].text_frame.text == "Israel Water Authority"            # untouched words follow the chart
    assert s2["credit"].text_frame.text == "Photo: JNF archive (hand)"      # hand words stay ...
    assert len(warnings) == 1 and "credit" in warnings[0]                   # ... and the change is reported
    assert s2["my note"].text_frame.text == "mine"
    assert [sum(sh.shape_type == 13 for sh in s.values()) for s in shapes(out)] == pics_before


# ---------------------------------------------------------------- the real show

AGREED = {   # talk slide -> label names that must be native text (agreed 2026-10-20, scope A + B)
    1: {"subtitle", "source"},
    2: {"who", "how", "who.1", "how.1", "who.5", "how.5"},
    3: {"credit"},
    5: {"row.1", "row.4", "models_start", "test_years"},
    6: {"left", "left.sub", "right", "right.sub"},
    7: {"row.1", "row.4"},
    8: {"legend.1", "legend.4", "inside"},
    9: {"lane.1", "lane.2", "footer"},
    10: {"star", "caption", "header.1", "model.1", "model.2"},
    11: {"credit"},
    12: {"stamp", "credit", "app", "below", "wet", "dry"},
}


@pytest.mark.skipif(not SHOW.exists(), reason="the show isn't built")
def test_the_show_has_the_agreed_native_text_and_still_runs_12_minutes():
    from pptx import Presentation
    prs = Presentation(SHOW)
    assert len(prs.slides) == 12
    for n, names in AGREED.items():
        have = {sh.name for sh in prs.slides[n - 1].shapes if sh.has_text_frame and sh.text_frame.text}
        missing = {m for m in names if not any(h == m or h.startswith(m + ".") for h in have)}
        assert not missing, f"slide {n}: no native text for {sorted(missing)}"
    secs = 0.0
    for sl in prs.slides:
        t = sl._element.xpath("./p:transition/@advTm")
        secs += int(t[0]) / 1000 if t else 0
    assert 11.5 * 60 <= secs <= 13 * 60
