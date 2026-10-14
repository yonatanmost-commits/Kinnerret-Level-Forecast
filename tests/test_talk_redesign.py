"""The deck's redesign from the Claude Design project "Kinneret Deck" (originator 2026-10-14):
the watercolour look is dropped; metaphor slides become real black-and-white photographs, and the
data-sources slide becomes the coverage chart with each agency's logo beside the records it supplies.

Contract fixed here, in docs/talk/:
- talk.md `object: photo:<path>` - a full-bleed photo slide. `text: <title> | <subtitle>` is set as
  live text (white) over the picture, and `credit: <words>` as a small credit line. The builder turns
  the picture black and white and fades its left side to dark, so the text reads.
- compose_coverage.py main() -> charts/coverage_slide.png: 1920 x 1080, the drawn coverage layer on
  the deck's plain surface (no painted backdrop), the three agency logos down the right edge.
- talk.md: a photo slide's notes never call the picture AI-made; the fountain picture is gone and
  the sources slide shows the coverage chart.
"""
import sys
from pathlib import Path

import pytest
from PIL import Image

ROOT = Path(__file__).resolve().parent.parent
TALK = ROOT / "docs" / "talk"
sys.path.insert(0, str(TALK))

LEDGER = {"freeze_date": {"value": "2026-10-11", "display": "11 October 2026", "unit": "",
                          "source": "fixture", "date": "2026-10-11", "kind": "freeze"}}


@pytest.fixture
def bd():
    import build_deck
    return build_deck


def colourful_photo(path: Path) -> Path:
    img = Image.new("RGB", (1500, 1000), (200, 40, 40))           # strongly red
    img.paste((40, 200, 60), (750, 0, 1500, 1000))                 # right half strongly green
    img.save(path)
    return path


def photo_slide(bd, tmp_path):
    from pptx import Presentation
    pic = colourful_photo(tmp_path / "statue.jpg")
    s = bd.Slide(id="centaur", object=f"photo:{pic}",
                 text="The Centaur Model | Expert judgement and AI, working together",
                 notes="A centaur.", extra={"credit": "Photo: Someone / Pexels"})
    out = tmp_path / "deck.pptx"
    warnings = bd.build_pptx([s], LEDGER, out=out)
    prs = Presentation(out)
    return prs, prs.slides[0], warnings


def test_photo_slide_is_full_bleed_and_black_and_white(bd, tmp_path):
    from io import BytesIO
    prs, slide, warnings = photo_slide(bd, tmp_path)
    assert not warnings
    pics = [sh for sh in slide.shapes if sh.shape_type == 13]
    assert len(pics) == 1
    pic = pics[0]
    assert (pic.left, pic.top, pic.width, pic.height) == (0, 0, prs.slide_width, prs.slide_height)
    img = Image.open(BytesIO(pic.image.blob)).convert("RGB")
    w, h = img.size
    for x in (w * 0.6, w * 0.8, w * 0.95):                          # the side the text doesn't sit on
        r, g, b = img.getpixel((int(x), h // 2))
        assert max(r, g, b) - min(r, g, b) <= 3, f"colour left in the picture at x={x}: {(r, g, b)}"
    # the left side is faded towards dark so white text reads
    assert sum(img.getpixel((5, h // 2))) < sum(img.getpixel((int(w * 0.95), h // 2)))


def test_photo_slide_title_and_subtitle_are_live_white_text(bd, tmp_path):
    prs, slide, _ = photo_slide(bd, tmp_path)
    runs = {r.text: r for sh in slide.shapes if sh.has_text_frame
            for p in sh.text_frame.paragraphs for r in p.runs}
    for want in ("The Centaur Model", "Expert judgement and AI, working together"):
        assert want in runs, f"missing live text {want!r}; found {list(runs)}"
        assert str(runs[want].font.color.rgb) == "FFFFFF"
    assert runs["The Centaur Model"].font.size.pt > runs[
        "Expert judgement and AI, working together"].font.size.pt
    assert "Photo: Someone / Pexels" in runs


def test_coverage_slide_is_the_chart_on_a_plain_surface_with_the_logos():
    import build_deck
    import compose_coverage
    compose_coverage.main()
    img = Image.open(TALK / "charts" / "coverage_slide.png").convert("RGB")
    assert img.size == (1920, 1080)
    surface = tuple(int(build_deck.SURFACE[i:i + 2], 16) for i in (0, 2, 4))
    for xy in [(10, 10), (1910, 10), (10, 1070), (1910, 1070), (960, 1060)]:  # corners + bottom strip
        assert all(abs(a - b) <= 2 for a, b in zip(img.getpixel(xy), surface)), \
            f"not the plain surface at {xy}: {img.getpixel(xy)}"
    logo_strip = img.crop((1715, 280, 1900, 860))
    inked = sum(1 for px in logo_strip.getdata() if sum(px) < 600)
    assert inked > 2000, "no logos down the right edge"


def test_script_photos_are_not_called_ai_made_and_the_fountain_is_gone(bd):
    slides = bd.parse_talk((TALK / "talk.md").read_text(encoding="utf-8"))
    by_id = {s.id: s for s in slides}
    assert by_id["centaur"].object.startswith("photo:")
    for s in slides:
        if s.object.startswith("photo:"):
            assert "AI-made" not in s.notes, f"slide {s.id}: a real photo is called AI-made"
            assert s.extra.get("credit"), f"slide {s.id}: photo without a credit"
    assert "sources.png" not in " ".join(s.object for s in slides)
    assert by_id["sources"].object == "chart:coverage_slide"
    assert "coverage" not in by_id                                  # merged into sources
