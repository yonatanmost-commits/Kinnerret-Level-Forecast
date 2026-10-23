"""
deck_diff.py - what was changed by hand in a deck, against a fresh build of the same source
(originator, 23 Oct: "ill ask you to run a diff, and fix"). Shapes are matched by name on the
master, the layouts and every slide: added, removed, moved or resized, re-worded or re-styled,
re-stacked; plus each slide's notes, and its animation and transition XML.

Run:  CLEAN_OUT=Kinneret_Talk_Submission_ref.pptx python docs/talk/build_clean.py
      python docs/talk/deck_diff.py docs/talk/Kinneret_Talk_Submission_ref.pptx docs/talk/Kinneret_Talk_Submission.pptx
Fold each change into its source (deck.html, the chart scripts, build_show.HAND_MOVES /
TITLE_BAR, build_clean), rebuild, and run it again: only intended differences should remain.
"""
import sys
from pptx import Presentation
from pptx.util import Emu
pt = lambda e: round(Emu(e).pt, 1) if e is not None else None
def desc(sh):
    d = {"type": str(sh.shape_type), "geom": (pt(sh.left), pt(sh.top), pt(sh.width), pt(sh.height))}
    try: d["rot"] = sh.rotation
    except Exception: pass
    if sh.has_text_frame:
        d["text"] = sh.text_frame.text
        runs = []
        for p in sh.text_frame.paragraphs:
            for r in p.runs:
                f = r.font
                c = None
                try: c = str(f.color.rgb) if f.color and f.color.type is not None else None
                except Exception: c = "theme"
                runs.append((f.name, f.size.pt if f.size else None, f.bold, c))
        d["runs"] = sorted(set(runs), key=str)
        d["align"] = [str(p.alignment) for p in sh.text_frame.paragraphs][:1]
    try:
        d["fill"] = str(sh.fill.fore_color.rgb) if sh.fill.type == 1 else str(sh.fill.type)
    except Exception: pass
    try: d["line"] = (str(sh.line.color.rgb) if sh.line.fill.type == 1 else str(sh.line.fill.type), pt(sh.line.width))
    except Exception: pass
    return d
def shapes(coll): 
    out = {}
    for i, sh in enumerate(coll):
        k = sh.name if sh.name not in out else f"{sh.name}#{i}"
        out[k] = (i, desc(sh))
    return out
def cmp(where, A, B):
    a, b = shapes(A), shapes(B)
    for k in a.keys() - b.keys(): print(f"{where}: REMOVED by hand: {k} {a[k][1]}")
    for k in b.keys() - a.keys(): print(f"{where}: ADDED by hand: {k} {b[k][1]}")
    for k in a.keys() & b.keys():
        da, db = a[k][1], b[k][1]
        for f in set(da) | set(db):
            if da.get(f) != db.get(f): print(f"{where}: {k}.{f}: {da.get(f)!r} -> {db.get(f)!r}")
    oa = [k for k in sorted(a, key=lambda k: a[k][0]) if k in b]; ob = [k for k in sorted(b, key=lambda k: b[k][0]) if k in a]
    if oa != ob: print(f"{where}: Z-ORDER {oa} -> {ob}")
ref, hand = Presentation(sys.argv[1]), Presentation(sys.argv[2])
print("size", ref.slide_width, hand.slide_width, ref.slide_height, hand.slide_height)
cmp("MASTER", ref.slide_master.shapes, hand.slide_master.shapes)
print("master bg xml differs:", ref.slide_master.element.find('.//{*}bg') is not None and
      __import__('lxml.etree',fromlist=['x']).tostring(ref.slide_master.element.find('.//{*}bg')) != __import__('lxml.etree',fromlist=['x']).tostring(hand.slide_master.element.find('.//{*}bg')))
for la, lb in zip(ref.slide_layouts, hand.slide_layouts): cmp(f"LAYOUT {la.name}->{lb.name}", la.shapes, lb.shapes)
print("layouts", [l.name for l in ref.slide_layouts], [l.name for l in hand.slide_layouts])
print("slides", len(ref.slides), len(hand.slides))
for i, (sa, sb) in enumerate(zip(ref.slides, hand.slides), 1):
    cmp(f"SLIDE {i}", sa.shapes, sb.shapes)
    na, nb = sa.notes_slide.notes_text_frame.text, sb.notes_slide.notes_text_frame.text
    if na != nb: print(f"SLIDE {i}: NOTES\n  ref: {na}\n  hand: {nb}")
    import re
    ta = re.findall(r"<p:(?:transition|timing)\b.*?</p:(?:transition|timing)>", sa.element.xml, re.S)
    tb = re.findall(r"<p:(?:transition|timing)\b.*?</p:(?:transition|timing)>", sb.element.xml, re.S)
    if ta != tb:
        ea = re.findall(r'presetClass="(\w+)" presetID="(\d+)"|nodeType="(\w+)"', "".join(ta)); eb = re.findall(r'presetClass="(\w+)" presetID="(\d+)"|nodeType="(\w+)"', "".join(tb))
        print(f"SLIDE {i}: TIMING/TRANSITION differs; effects ref={len(ea)} hand={len(eb)}", "" if ea == eb else f"\n  ref {ea}\n  hand {eb}")
        tra = re.findall(r"<p:transition[^>]*>", "".join(ta)); trb = re.findall(r"<p:transition[^>]*>", "".join(tb))
        if tra != trb: print("  transition", tra, "->", trb)
