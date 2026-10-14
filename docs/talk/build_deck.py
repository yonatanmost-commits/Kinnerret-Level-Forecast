"""
build_deck.py - talk.md + numbers ledger + charts + pictures -> talk.pptx and talk.pdf.

Gates that run on every build (plan D6, D8; tests/test_talk_gates.py):
  check_stamp   the ledger must match the hash the truth tests last passed on (truth_stamp.json);
                a match means the facts are already proven, so the truth tests are not re-run
  check_numbers no digit may be typed into the text: numbers enter only as {ledger.key}
                placeholders, or as a phrase on the file's `allow:` line

talk.md format:
  allow: 2023 | 15 minutes            exempt phrases, " | "-separated (one line, near the top)

  ## slide <id>
  object: chart:<name> | picture:<path> | photo:<path> | table:<name>
  credit: <words>                     a photo's credit line, set small in its corner
  qr: app                             optional: the app's QR code in the corner (plan D18)
  text: <on-slide words>
  notes:
  <speaker notes, any number of lines, with {placeholders}>
  accepted: <short content hash, or empty>
"""
from __future__ import annotations

import hashlib
import json
import re
from dataclasses import dataclass, field
from pathlib import Path

HERE = Path(__file__).resolve().parent
TALK_MD = HERE / "talk.md"
QA_MD = HERE / "qa_prep.md"
LEDGER_PATH = HERE / "numbers_ledger.json"
STAMP_PATH = HERE / "truth_stamp.json"

PLACEHOLDER = re.compile(r"\{([A-Za-z0-9_.]+)\}")
NUMBER = re.compile(r"\d(?:[\d.,:/]*\d)?")
STRUCTURE = ("allow:", "object:", "accepted:", "## slide")


class BuildStopped(Exception):
    """A gate refused the build; the message says why."""


# ─── gates ──────────────────────────────────────────────────────────────────

def check_stamp(ledger_path: Path = LEDGER_PATH, stamp_path: Path = STAMP_PATH) -> None:
    ledger_path, stamp_path = Path(ledger_path), Path(stamp_path)
    if not stamp_path.exists():
        raise BuildStopped(f"no truth stamp ({stamp_path.name}): the truth tests have not passed "
                           f"on these facts - run Automation/talk_freeze_run.py")
    stamped = json.loads(stamp_path.read_text(encoding="utf-8")).get("ledger_sha256")
    current = hashlib.sha256(ledger_path.read_bytes()).hexdigest()
    if stamped != current:
        raise BuildStopped("the numbers ledger changed since the truth tests last passed - run "
                           "Automation/talk_freeze_run.py so the truth tests pass first")


def allow_list(text: str) -> list[str]:
    for line in text.splitlines():
        if line.startswith("allow:"):
            return [p.strip() for p in line[len("allow:"):].split("|") if p.strip()]
    return []


def typed_numbers(text: str) -> list[tuple[int, str]]:
    """(line number, token) for every digit run outside placeholders and allowed phrases."""
    allowed = sorted(allow_list(text), key=len, reverse=True)
    found = []
    for i, line in enumerate(text.splitlines(), 1):
        if line.lstrip().startswith(STRUCTURE):
            continue
        scan = PLACEHOLDER.sub(" ", line)
        for phrase in allowed:
            scan = re.sub(rf"(?<![\w.]){re.escape(phrase)}(?![\w]|\.\d)", " ", scan)
        found += [(i, m.group()) for m in NUMBER.finditer(scan)]
    return found


def check_numbers(text_path: Path, ledger: dict) -> None:
    text_path = Path(text_path)
    text = text_path.read_text(encoding="utf-8")
    problems = [f"line {i}: typed number '{tok}'" for i, tok in typed_numbers(text)]
    unknown = sorted({k for k in PLACEHOLDER.findall(text) if k not in ledger})
    problems += [f"unknown placeholder {{{k}}}" for k in unknown]
    if problems:
        raise BuildStopped(f"{text_path.name}: numbers must come from the ledger -\n  "
                           + "\n  ".join(problems))


# ─── talk.md ────────────────────────────────────────────────────────────────

@dataclass
class Slide:
    id: str
    object: str = ""
    text: str = ""
    notes: str = ""
    accepted: str = ""
    extra: dict = field(default_factory=dict)

    def content_hash(self) -> str:
        body = "\n".join([self.id, self.object, self.text, self.notes])
        return hashlib.sha256(body.encode("utf-8")).hexdigest()[:8]


def parse_talk(text: str) -> list[Slide]:
    slides: list[Slide] = []
    cur, in_notes = None, False
    for line in text.splitlines():
        if line.startswith("## slide"):
            cur = Slide(id=line[len("## slide"):].strip())
            slides.append(cur)
            in_notes = False
            continue
        if cur is None:
            continue
        if line.startswith("accepted:"):
            cur.accepted = line[len("accepted:"):].strip()
            in_notes = False
        elif in_notes:
            cur.notes += line + "\n"
        elif line.startswith("notes:"):
            in_notes = True
        elif line.startswith("credit:"):
            cur.extra["credit"] = line[len("credit:"):].strip()
        elif line.startswith("qr:"):
            cur.extra["qr"] = line[len("qr:"):].strip()
        elif line.startswith("object:"):
            cur.object = line[len("object:"):].strip()
        elif line.startswith("text:"):
            cur.text = line[len("text:"):].strip()
    for s in slides:
        s.notes = s.notes.strip()
    return slides


def fill(text: str, ledger: dict) -> str:
    return PLACEHOLDER.sub(lambda m: str(ledger[m.group(1)]["display"]), text)


# ─── numbers_check.md (spec A11, A14) ───────────────────────────────────────

# Every phrase on talk.md's allow: line needs a source here, so Adnan can check it (plan D18).
ALLOW_SOURCES = {
    "2023": "the year the lake's rain gauge went offline and the pipeline turned it into zero rain; "
            "fixed 2026-05-30 (R² history step ac2ede5)",
    "2005": "the PAL/CSS Freestyle chess tournament, won by two amateurs using three computers: "
            "G. Kasparov, 'The Chess Master and the Computer', New York Review of Books, 11 Feb 2010",
    "1960": "the start of the ERA5 weather reanalysis used to rebuild the past (Open-Meteo archive, "
            "1960-01-01 onward; commit ea3a1af; paper Table C)",
}

CROSSING_BANDS = [("0_7", "1-7"), ("7_14", "8-14"), ("14_30", "15-30"), ("30_60", "31-60"),
                  ("60_90", "61-90"), ("90_120", "91-120"), ("120_plus", "over 120")]


def write_numbers_check(ledger: dict, out: Path = HERE / "numbers_check.md") -> Path:
    """Every number the talk uses, where it appears, its value and its date - for the originator
    and Adnan to read. Generated; never edited by hand."""
    uses: dict[str, list[str]] = {}
    for path in (TALK_MD, QA_MD):
        if not path.exists():
            continue
        text = path.read_text(encoding="utf-8")
        for s in parse_talk(text) if path == TALK_MD else []:
            for k in PLACEHOLDER.findall(s.text + "\n" + s.notes):
                uses.setdefault(k, []).append(f"slide {s.id}")
        if path == QA_MD:
            for k in PLACEHOLDER.findall(text):
                uses.setdefault(k, []).append("Q&A sheet")
    for s in parse_talk(TALK_MD.read_text(encoding="utf-8")):
        if s.object.startswith("chart:"):
            uses.setdefault(f"(chart) {s.object[6:]}", []).append(f"slide {s.id}")

    fd = ledger["freeze_date"]["value"]
    lines = ["# Numbers check: Kinneret Conference talk", "",
             f"Generated by build_deck.py from numbers_ledger.json. Data freeze: **{fd}**.",
             "Measured numbers come from one re-run on the freeze. Historical numbers keep the date "
             "they were first measured and show it wherever they appear.", "",
             "| Number | Shown as | Unit | Where | Kind | Date | Source |", "|---|---|---|---|---|---|---|"]
    for k in sorted(k for k in uses if not k.startswith("(chart)")):
        e = ledger[k]
        date = e["date"] + (f" ({e['measured_note']})" if e.get("measured_note") else "")
        where = ", ".join(dict.fromkeys(uses[k]))
        lines.append(f"| `{k}` | {e['display']} | {e['unit']} | "
                     f"{where} | {e['kind']} | {date} | {e['source']} |")
    charts = sorted(k for k in uses if k.startswith("(chart)"))
    lines += ["", "Charts draw only from the ledger (charts.py chart_data); "
              f"charts used: {', '.join(c[8:] for c in charts)}.", ""]

    allowed = allow_list(TALK_MD.read_text(encoding="utf-8"))
    if allowed:
        lines += ["## Typed phrases allowed in the script", "", "| Phrase | Source |", "|---|---|"]
        lines += [f"| {a} | {ALLOW_SOURCES.get(a, '**no source given**')} |" for a in allowed]
        lines.append("")

    if "crossing.issued" in ledger:
        lines += ["## Red-line crossing hindcast, per band of days of warning", "",
                  f"{ledger['crossing.n']['value']} crossings of -213.0 m since "
                  f"{ledger['crossing.first_year']['value']}; a forecast issued every day from the last "
                  f"rain of at least {ledger['crossing.last_rain_mm']['value']} mm. 'Paired' = only the "
                  "forecasts both methods brought to the line; every spoken comparison uses paired.", "",
                  "| Days of warning | Issued | Scaled never crossed | Plain never crossed | "
                  "Paired n | Paired MAE scaled (days) | Paired MAE plain (days) |",
                  "|---|---|---|---|---|---|---|"]
        for b, label in CROSSING_BANDS:
            g = lambda k: ledger[k]["value"]
            lines.append(f"| {label} | {g(f'crossing.issued_{b}')} | {g(f'crossing.scaled.never_{b}')} | "
                         f"{g(f'crossing.unscaled.never_{b}')} | {g(f'crossing.paired.n_{b}')} | "
                         f"{g(f'crossing.paired.scaled.mae_{b}')} | {g(f'crossing.paired.unscaled.mae_{b}')} |")
        lines.append("")
    out.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return out


# ─── the deck (plan D1, D2, D10) ────────────────────────────────────────────
# One object per slide on a warm near-white page. Chart slides show the chart alone: its headline
# is the slide's point. Metaphor pictures run full-bleed and the prompts leave their left third
# empty, so any slide text sits there. A picture not made yet becomes a labelled placeholder and a
# build warning, never a silent gap.

PPTX_PATH = HERE / "talk.pptx"
PDF_PATH = HERE / "talk.pdf"
CHARTS_DIR = HERE / "charts"
SLIDE_W, SLIDE_H = 13.333, 7.5                 # inches, 16:9
SURFACE, INK, INK_2, MUTED = "fbfaf7", "1d1d1b", "5f5d57", "e9e6de"
# One family at weights heavier than Light (originator, 2026-10-13); the weight is in the font
# name, so text is never set bold on top of it.
HEAD_FONT, BODY_FONT = "Segoe UI Semibold", "Segoe UI Semilight"
QR_IMAGE = HERE / "images" / "qr.png"


def _rgb(hex6):
    from pptx.dml.color import RGBColor
    return RGBColor.from_string(hex6)


def _text(slide, text, left, top, width, height, size, font=HEAD_FONT, bold=False,
          color=INK, align="left", anchor="middle"):
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Inches, Pt
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE}[anchor]
    for i, line in enumerate(text.split("\n")):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[align]
        run = para.add_run()
        run.text = line
        run.font.size, run.font.name, run.font.bold = Pt(size), font, bold
        run.font.color.rgb = _rgb(color)
    return box


def _picture_fit(slide, path, left, top, width, height):
    """The whole picture, as large as fits in the box, centred in it."""
    from PIL import Image
    from pptx.util import Inches
    w, h = Image.open(path).size
    scale = min(width / w, height / h)
    pw, ph = w * scale, h * scale
    return slide.shapes.add_picture(str(path), Inches(left + (width - pw) / 2),
                                    Inches(top + (height - ph) / 2), Inches(pw), Inches(ph))


def _picture_cover(slide, path):
    """Full-bleed: the picture fills the slide, cropped evenly on the long side."""
    from PIL import Image
    from pptx.util import Inches
    w, h = Image.open(path).size
    pic = slide.shapes.add_picture(str(path), 0, 0, Inches(SLIDE_W), Inches(SLIDE_H))
    excess = 1 - (SLIDE_W / SLIDE_H) / (w / h)
    if excess > 0:                               # picture wider than the slide
        pic.crop_left = pic.crop_right = excess / 2
    elif excess < 0:                             # taller
        cut = 1 - (w / h) / (SLIDE_W / SLIDE_H)
        pic.crop_top = pic.crop_bottom = cut / 2
    return pic


# A photo slide (Claude Design "Kinneret Deck", 2026-10-14): the photograph in black and white,
# full-bleed, its left side faded to a dark slate so the white title and subtitle read there.
FADE_RGB, FADE_ALPHA, FADE_SHARE = (10, 20, 24), 0.88, 1000 / 1920
ACCENT = "009c8f"                              # the deck's teal (charts.py TEAL)
PHOTO_SUB = "FFFFFF"


def _photo_image(path) -> "BytesIO":
    """The picture cropped to the slide's shape, black and white, faded on the left; as JPEG."""
    from io import BytesIO
    from PIL import Image, ImageOps
    img = ImageOps.fit(Image.open(path).convert("L"), (1920, 1080), Image.LANCZOS).convert("RGB")
    fade_w = round(1920 * FADE_SHARE)
    ramp = ImageOps.mirror(Image.linear_gradient("L").rotate(90, expand=True)).resize(
        (fade_w, 1080))                                  # 255 at the left edge -> 0
    mask = Image.new("L", img.size, 0)
    mask.paste(ramp.point(lambda v: round(v * FADE_ALPHA)), (0, 0))
    img.paste(Image.new("RGB", img.size, FADE_RGB), (0, 0), mask)
    buf = BytesIO()
    img.save(buf, "JPEG", quality=90)
    buf.seek(0)
    return buf


def _photo_slide(slide, path, text, credit):
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches
    slide.shapes.add_picture(_photo_image(path), 0, 0, Inches(SLIDE_W), Inches(SLIDE_H))
    title, _, sub = (t.strip() for t in text.partition("|"))
    px = SLIDE_W / 1920                                      # the design's pixels, in inches
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(120 * px), Inches(690 * px),
                                 Inches(96 * px), Inches(6 * px))
    bar.fill.solid()
    bar.fill.fore_color.rgb = _rgb(ACCENT)
    bar.line.fill.background()
    for box in (_text(slide, title, 120 * px, 716 * px, 1500 * px, 140 * px, 54, color="FFFFFF",
                      anchor="top"),
                _text(slide, sub, 120 * px, 872 * px, 1500 * px, 70 * px, 28, font=BODY_FONT,
                      color=PHOTO_SUB, anchor="top") if sub else None,
                _text(slide, credit, 120 * px, 1020 * px, 900 * px, 36 * px, 10, font=BODY_FONT,
                      color="b8c0c2") if credit else None):
        if box is not None:
            box.text_frame.margin_left = box.text_frame.margin_right = 0


def _placeholder(slide, name):
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches
    box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(4.9), Inches(0.9), Inches(7.6),
                                 Inches(5.7))
    box.fill.solid()
    box.fill.fore_color.rgb = _rgb(MUTED)
    box.line.fill.background()
    _text(slide, f"picture to make by hand:\n{name}.png\n(prompt in images/{name}.prompt.txt)",
          4.9, 0.9, 7.6, 5.7, 24, font=BODY_FONT, color=INK_2, align="center")


def _lay_out(slide, s: Slide, text: str, warnings: list[str]):
    kind, _, target = s.object.partition(":")
    if kind == "chart":
        path = CHARTS_DIR / f"{target}.png"
        if not path.exists():
            warnings.append(f"slide {s.id}: chart {path.name} missing (run charts.py)")
            return
        _picture_fit(slide, path, 0, 0, SLIDE_W, SLIDE_H)
    elif s.id == "title":
        title, _, names = text.partition("|")
        _picture_fit(slide, HERE / target, 7.6, 0.4, 5.2, 6.7)
        _text(slide, title.strip(), 0.8, 1.6, 6.6, 2.6, 48, anchor="top")
        _text(slide, names.strip(), 0.8, 4.4, 6.6, 1.0, 28, font=BODY_FONT, color=INK_2,
              anchor="top")
    elif target == "images/qr.png":
        _picture_fit(slide, HERE / target, 7.9, 1.25, 5.0, 5.0)
        _text(slide, text.replace(" · ", "\n"), 0.8, 1.25, 6.6, 5.0, 40)
    elif kind == "photo":
        _photo_slide(slide, HERE / target, text, s.extra.get("credit", ""))
    elif kind == "picture":
        path = HERE / target
        if path.exists():
            _picture_cover(slide, path)
        else:
            _placeholder(slide, path.stem)
            warnings.append(f"slide {s.id}: picture {target} not made yet")
        if text:
            _text(slide, text, 0.6, 0.8, 4.0, 5.9, 36)
    else:
        warnings.append(f"slide {s.id}: unknown object {s.object!r}")
    if s.extra.get("qr") == "app":
        _picture_fit(slide, QR_IMAGE, SLIDE_W - 1.75, 0.3, 1.45, 1.45)     # top right: clear of the data


def build_pptx(slides: list[Slide], ledger: dict, out: Path = PPTX_PATH) -> list[str]:
    """talk.md's slides -> talk.pptx, with filled speaker notes. Returns the build warnings."""
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(SLIDE_W), Inches(SLIDE_H)
    blank = prs.slide_layouts[6]
    warnings: list[str] = []
    for s in slides:
        slide = prs.slides.add_slide(blank)
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = _rgb(SURFACE)
        _lay_out(slide, s, fill(s.text, ledger), warnings)
        slide.notes_slide.notes_text_frame.text = fill(s.notes, ledger)
        if len(fill(s.text, ledger).replace("|", " ").split()) > 12:
            warnings.append(f"slide {s.id}: more than 12 words of slide text")
    try:
        prs.save(out)
    except PermissionError:
        raise BuildStopped(f"{out.name} is open in PowerPoint - close it and build again")
    return warnings


# ─── PowerPoint: the PDF and the slide pictures (plan D2) ───────────────────

def export(pptx: Path = PPTX_PATH, pdf: Path | None = PDF_PATH,
           pictures: list[int] | None = None) -> list[Path]:
    """PowerPoint itself renders the PDF and slide pictures (slide_NN.png beside this file):
    the real renderer is the proof for fonts and layout. Leaves a running PowerPoint as it was."""
    import win32com.client
    if pdf is not None and Path(pdf).exists():
        try:
            Path(pdf).open("r+b").close()
        except PermissionError:
            raise BuildStopped(f"{Path(pdf).name} is open in a PDF viewer - close it and build again")
    app = win32com.client.Dispatch("PowerPoint.Application")
    was_open = app.Presentations.Count
    pres = app.Presentations.Open(str(pptx), True, False, False)   # read-only, no window
    made = []
    try:
        if pdf is not None:
            pres.SaveAs(str(pdf), 32)                                   # ppSaveAsPDF
        for n in pictures or []:
            path = HERE / f"slide_{n:02d}.png"
            pres.Slides(n).Export(str(path), "PNG", 1920, 1080)
            made.append(path)
    finally:
        pres.Close()
        if was_open == 0 and app.Presentations.Count == 0:
            app.Quit()
    return made


def check_render(slides: list[Slide], pdf: Path = PDF_PATH) -> list[str]:
    """Build warnings A9: the PDF has the deck's slides in order, the fonts are embedded, and
    every chart or picture slide carries an image."""
    import fitz
    warnings = []
    doc = fitz.open(pdf)
    if doc.page_count != len(slides):
        warnings.append(f"PDF has {doc.page_count} pages, the deck {len(slides)} slides")
    fonts = {f[3].split("+")[-1] for page in doc for f in page.get_fonts()}
    squash = lambda name: re.sub(r"[\s-]", "", name).lower()
    for want in (HEAD_FONT, BODY_FONT):
        if not any(squash(want) in squash(f) for f in fonts):
            warnings.append(f"font {want} not embedded in the PDF (fonts: {sorted(fonts)})")
    for i, (s, page) in enumerate(zip(slides, doc), 1):
        if s.object and not page.get_images():
            warnings.append(f"page {i} ({s.id}): no image rendered")
        drawn = "" if s.object.startswith("chart:") else s.text    # chart slides show the chart only
        words = re.findall(r"[A-Za-z]{4,}", drawn)[:3]
        text = page.get_text()
        if words and not all(w in text for w in words):
            warnings.append(f"page {i}: expected slide {s.id}'s text, found {text[:40]!r}")
    return warnings


def unaccepted(slides: list[Slide]) -> list[str]:
    """A12: slides never accepted, or changed since they were."""
    return [f"{s.id} ({'changed since accepted' if s.accepted else 'not accepted yet'}, "
            f"hash {s.content_hash()})" for s in slides if s.accepted != s.content_hash()]


def build(slide: int | None = None) -> None:
    ledger = json.loads(LEDGER_PATH.read_text(encoding="utf-8"))
    check_stamp()
    for path in (TALK_MD, QA_MD):
        check_numbers(path, ledger)
    slides = parse_talk(TALK_MD.read_text(encoding="utf-8"))
    warnings = build_pptx(slides, ledger)
    write_numbers_check(ledger)
    if slide is not None:
        if not 1 <= slide <= len(slides):
            raise BuildStopped(f"no slide {slide}; the deck has {len(slides)}")
        made = export(pdf=None, pictures=[slide])
        print(f"slide {slide} ({slides[slide - 1].id}): {made[0].relative_to(HERE)}")
    else:
        export(pictures=list(range(1, len(slides) + 1)))
        warnings += check_render(slides)
        print(f"built {PPTX_PATH.name} and {PDF_PATH.name}: {len(slides)} slides, "
              f"pictures slide_01..{len(slides):02d}.png")
    for w in warnings:
        print("WARNING", w)
    todo = unaccepted(slides)
    if todo:
        print(f"{len(todo)} of {len(slides)} slides not accepted:", "; ".join(todo))


def main(argv=None) -> int:
    import argparse
    import sys
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    ap = argparse.ArgumentParser(description="talk.md -> talk.pptx, talk.pdf, slide pictures")
    ap.add_argument("--slide", type=int, help="build, then render only slide N as a picture")
    args = ap.parse_args(argv)
    try:
        build(args.slide)
    except BuildStopped as e:
        print("BUILD STOPPED:", e)
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
